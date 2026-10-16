"""Strict parsing and safe artifact creation for BUSY Excel exports.

Office Hub does not inspect BUSY's Access database. It accepts an explicit workbook contract so a
bad or changed export fails loudly instead of producing plausible-looking financial figures.
"""

from __future__ import annotations

import re
import uuid
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


class BusyReportError(ValueError):
    """Raised when a workbook is not a recognized, safe report export."""


SUPPORTED_REPORTS = frozenset({"profit_loss", "balance_sheet"})
REQUIRED_HEADERS = frozenset({"section", "account"})
AMOUNT_HEADERS = frozenset({"amount", "debit", "credit"})


@dataclass(frozen=True)
class ReportRow:
    section: str
    account: str
    debit: float
    credit: float
    amount: float


@dataclass(frozen=True)
class ParsedBusyReport:
    report_type: str
    source_file: Path
    rows: tuple[ReportRow, ...]
    totals: dict[str, float]
    warnings: tuple[str, ...]
    review_required: bool


def normalize_header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(value or "").strip().lower()).strip("_")


def parse_amount(value: Any) -> float:
    if value is None or value == "":
        return 0.0
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace(",", "").replace("₹", "")
    if text in {"-", "—", "–"}:
        return 0.0
    negative = text.startswith("(") and text.endswith(")")
    if negative:
        text = text[1:-1]
    try:
        parsed = float(text)
    except ValueError as exc:
        raise BusyReportError(f"Invalid amount: {value!r}") from exc
    return -parsed if negative else parsed


def _safe_input_path(path: Path, allowed_root: Path) -> Path:
    resolved = path.expanduser().resolve()
    root = allowed_root.expanduser().resolve()
    if not resolved.is_relative_to(root):
        raise BusyReportError("Workbook must be inside the configured BUSY inbox")
    if resolved.suffix.lower() != ".xlsx":
        raise BusyReportError("Only .xlsx exports are supported; .xls and .xlsm are rejected")
    if not resolved.is_file():
        raise FileNotFoundError(f"BUSY workbook does not exist: {resolved}")
    return resolved


def parse_busy_workbook(path: Path, report_type: str, allowed_root: Path) -> ParsedBusyReport:
    """Parse the first sheet using the explicit Office Hub export contract."""

    if report_type not in SUPPORTED_REPORTS:
        raise BusyReportError(f"Unsupported report type: {report_type}")
    source = _safe_input_path(path, allowed_root)
    workbook = load_workbook(source, read_only=True, data_only=True)
    try:
        if not workbook.worksheets:
            raise BusyReportError("Workbook has no worksheets")
        sheet = workbook.worksheets[0]
        values = sheet.iter_rows(values_only=True)
        try:
            raw_headers = next(values)
        except StopIteration as exc:
            raise BusyReportError("Workbook is empty") from exc
        headers = [normalize_header(value) for value in raw_headers]
        positions = {header: index for index, header in enumerate(headers) if header}
        duplicate_headers = {
            header for header in headers if header and headers.count(header) > 1
        }
        if duplicate_headers:
            raise BusyReportError(
                f"Workbook contains duplicate headers: {', '.join(sorted(duplicate_headers))}"
            )
        missing = REQUIRED_HEADERS - positions.keys()
        if missing or not AMOUNT_HEADERS.intersection(positions):
            raise BusyReportError(
                "Workbook must contain Section, Account, and at least one of Amount, Debit, Credit"
            )

        rows: list[ReportRow] = []
        for raw_row in values:
            section = str(raw_row[positions["section"]] or "").strip()
            account = str(raw_row[positions["account"]] or "").strip()
            if not section and not account:
                continue
            debit = _cell_amount(raw_row, positions, "debit")
            credit = _cell_amount(raw_row, positions, "credit")
            amount = _cell_amount(raw_row, positions, "amount")
            if "amount" not in positions:
                amount = credit - debit
            rows.append(ReportRow(section, account, debit, credit, amount))
    finally:
        workbook.close()

    if not rows:
        raise BusyReportError("Workbook contains no report rows")
    totals, warnings = _calculate_totals(report_type, rows)
    return ParsedBusyReport(
        report_type=report_type,
        source_file=source,
        rows=tuple(rows),
        totals=totals,
        warnings=tuple(warnings),
        review_required=bool(warnings),
    )


def _cell_amount(row: tuple[Any, ...], positions: dict[str, int], header: str) -> float:
    position = positions.get(header)
    return parse_amount(row[position]) if position is not None and position < len(row) else 0.0


def _calculate_totals(
    report_type: str,
    rows: list[ReportRow],
) -> tuple[dict[str, float], list[str]]:
    grouped: dict[str, float] = {}
    for row in rows:
        key = row.section.casefold()
        grouped[key] = grouped.get(key, 0.0) + row.amount

    warnings: list[str] = []
    if report_type == "profit_loss":
        revenue = sum(
            value
            for key, value in grouped.items()
            if _matches(key, "revenue", "income", "sales")
        )
        expenses = sum(
            value
            for key, value in grouped.items()
            if _matches(key, "expense", "cost", "purchase")
        )
        if revenue == 0 and expenses == 0:
            warnings.append("No revenue/income or expense/cost sections were recognized")
        return {
            "revenue": revenue,
            "expenses": expenses,
            "net_profit": revenue - expenses,
        }, warnings

    assets = sum(value for key, value in grouped.items() if _matches(key, "asset"))
    liabilities = sum(value for key, value in grouped.items() if _matches(key, "liabil"))
    equity = sum(value for key, value in grouped.items() if _matches(key, "equity", "capital"))
    difference = assets - (liabilities + equity)
    if abs(difference) > 0.01:
        warnings.append(f"Balance sheet does not reconcile; difference={difference:.2f}")
    return {
        "assets": assets,
        "liabilities": liabilities,
        "equity": equity,
        "balance_difference": difference,
    }, warnings


def _matches(value: str, *needles: str) -> bool:
    return any(needle in value for needle in needles)


def write_normalized_workbook(report: ParsedBusyReport, output_dir: Path) -> Path:
    """Write a sanitized, unique normalized workbook using an atomic replace."""

    output_dir.mkdir(parents=True, exist_ok=True)
    output = output_dir / f"{report.report_type}-{uuid.uuid4().hex}.xlsx"
    temporary = output.with_suffix(".tmp.xlsx")
    workbook = Workbook()
    summary = workbook.active
    if summary is None:
        raise BusyReportError("Unable to create normalized workbook summary sheet")
    summary.title = "Summary"
    summary.append(["Report", report.report_type])
    summary.append(["Source file", _safe_excel_text(report.source_file.name)])
    summary.append([])
    summary.append(["Metric", "Value"])
    for key, value in report.totals.items():
        summary.append([key, value])
    summary.append([])
    summary.append(["Review required", report.review_required])
    if report.warnings:
        summary.append(["Warnings", _safe_excel_text(" | ".join(report.warnings))])

    rows_sheet = workbook.create_sheet("Rows")
    rows_sheet.append(["Section", "Account", "Debit", "Credit", "Amount"])
    for row in report.rows:
        rows_sheet.append(
            [
                _safe_excel_text(row.section),
                _safe_excel_text(row.account),
                row.debit,
                row.credit,
                row.amount,
            ]
        )
    workbook.save(temporary)
    temporary.replace(output)
    return output


def _safe_excel_text(value: str) -> str:
    """Prevent imported text from being interpreted as an Excel formula."""

    if value and value[0] in "=+-@":
        return "'" + value
    return value
